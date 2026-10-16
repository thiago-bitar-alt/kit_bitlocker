# -*- coding: utf-8 -*-
"""Gera TUTORIAL_BitLocker.docx (tutorial do kit para o setor de forense). Uso: python gera_tutorial.py"""
import os

import docx
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

AQUI = os.path.dirname(os.path.abspath(__file__))
d = docx.Document()
for s in d.sections:
    s.left_margin = s.right_margin = Cm(2.2)
    s.top_margin = s.bottom_margin = Cm(2)
d.styles["Normal"].font.name = "Arial"
d.styles["Normal"].font.size = Pt(11)


conta = [0]  # numeração dos itens, reiniciada a cada seção


def h(t, n=1):
    conta[0] = 0
    d.add_heading(t, n)


def p(t, negrito=False):
    par = d.add_paragraph()
    par.add_run(t).bold = negrito
    return par


def item(t, numerado=False):
    if not numerado:
        d.add_paragraph(t, style="List Bullet")
        return
    conta[0] += 1
    par = d.add_paragraph(f"{conta[0]}. {t}")
    par.paragraph_format.left_indent = Cm(0.6)
    par.paragraph_format.first_line_indent = Cm(-0.5)


def cmd(t):
    par = d.add_paragraph()
    par.paragraph_format.left_indent = Cm(0.6)
    r = par.add_run(t)
    r.font.name = "Consolas"
    r.font.size = Pt(9.5)
    sombra = OxmlElement("w:shd")
    sombra.set(qn("w:val"), "clear")
    sombra.set(qn("w:fill"), "EDEDED")
    par._p.get_or_add_pPr().append(sombra)


def aviso(t):
    par = d.add_paragraph()
    r = par.add_run("Atenção: " + t)
    r.bold = True
    r.font.color.rgb = RGBColor(0xA0, 0x10, 0x10)


def tabela(cab, linhas, larg):
    t = d.add_table(rows=1, cols=len(cab))
    t.style = "Table Grid"
    for c, txt in zip(t.rows[0].cells, cab):
        c.text = txt
        c.paragraphs[0].runs[0].bold = True
    for l in linhas:
        for c, txt in zip(t.add_row().cells, l):
            c.text = txt
    for row in t.rows:
        for c, w in zip(row.cells, larg):
            c.width = Cm(w)
    d.add_paragraph()


tit = d.add_heading("Decifragem de BitLocker em imagens de notebooks (Windows 10 e 11)", 0)
p("Tutorial do kit_bitlocker: do zero até a imagem pronta para o IPED")

h("1. Quando usar")
p("O Windows 11 liga sozinho a criptografia do disco (BitLocker / \"Criptografia do dispositivo\") na maioria "
  "dos notebooks atuais. O IPED não decifra BitLocker: processa a partição cifrada como espaço bruto. Sinais de que "
  "a imagem tem BitLocker:")
item("o relatório do equipamento de extração (ex.: Logicube, \"Source Partition Information\") mostra "
     "Encryption=\"Bitlocker\";")
item("no IPED aparecem milhares de itens \"Não Alocado\"/\"Imagens de Disco\" e nenhuma pasta Windows ou Users;")
item("na visualização hexadecimal do início da partição aparece o texto -FVE-FS-.")
p("O kit diz se a chave está guardada \"em claro\" no próprio disco (decifra sem senha) ou se é preciso a chave de "
  "recuperação, e gera uma NOVA imagem com o Windows decifrado, sem alterar a original.")

h("2. O que é preciso")
item("Computador com Windows 10/11 de 64 bits (não precisa ser o da extração).")
item("Python 3.12 ou 3.13, 64 bits (PASSO 1).")
item("A pasta kit_bitlocker (scripts + pasta \"pacotes\", que permite instalar sem internet).")
item("Espaço livre, num disco diferente do da evidência, igual ao tamanho da imagem (ex.: notebook de 512 GB = "
     "~477 GB livres). De preferência SSD interno: ler e gravar no mesmo HD USB deixa o processo muito lento.")
item("A imagem bruta do disco (.dd ou .001; se dividida em .001, .002..., aponte para o .001). Para .E01, ver item 10.")

h("3. PASSO 0 - Na apreensão (se o notebook estiver LIGADO e desbloqueado)")
p("Desligado, a chave pode ficar presa ao chip TPM do notebook e a imagem sozinha não basta. Se o equipamento "
  "estiver ligado e desbloqueado:")
item("não desligue;", True)
item("abra o Prompt de Comando como administrador e execute o comando abaixo; fotografe a tela (a \"Senha "
     "Numérica\" de 48 dígitos é a chave de recuperação):", True)
cmd("manage-bde -protectors -get C:")
item("se possível, faça também a coleta com o sistema ligado, conforme o procedimento do setor.", True)

h("4. PASSO 1 - Instalar o Python (uma vez por computador)")
item("Baixe o instalador do Python 3.13 (Windows installer 64-bit) em https://www.python.org/downloads/windows/ "
     "(em máquina sem internet, leve o instalador num pendrive).", True)
item("Execute o instalador e marque a opção \"Add python.exe to PATH\" antes de clicar em \"Install Now\".", True)
item("Confira: abra o Prompt de Comando e digite o comando abaixo; deve aparecer \"Python 3.13...\".", True)
cmd("python --version")

h("5. PASSO 2 - Instalar o kit (uma vez por computador)")
item("Copie a pasta kit_bitlocker para o computador, por exemplo C:\\Ferramentas\\kit_bitlocker.", True)
item("Dê dois cliques em instalar.bat. Ele cria um ambiente próprio (.venv) e instala os pacotes da pasta "
     "\"pacotes\", sem internet. Ao final deve aparecer \"[OK] Kit instalado\".", True)

h("6. PASSO 3 - Checar a imagem (1 minuto, só leitura)")
p("Abra o Prompt de Comando na pasta do kit (na barra de endereço do Explorer, digite cmd e Enter) e execute, "
  "trocando o caminho pelo da imagem:")
cmd('checar.bat "D:\\Caso\\Evidencia\\DDCapture.001"')
p("O kit lista as partições do disco. Para a partição marcada como BITLOCKER:")
tabela(["Resultado", "O que fazer"],
       [["Protetores: CHAVE EM CLARO", "Decifra sem senha: PASSO 4."],
        ["Outros protetores (ou \"provável TPM\")",
         "A imagem sozinha não basta: PASSO 5 (anotar o ID do protetor)."],
        ["Nenhuma partição com BitLocker", "Processar a imagem original no IPED normalmente."]], [6.5, 10])

h("7. PASSO 4 - Decifrar (chave em claro)")
p("No Prompt, na pasta do kit, informe a imagem, a pasta de saída e o SHA-256 que consta no relatório da extração "
  "(para o kit conferir se a imagem lida é a mesma da extração):")
cmd('decifrar.bat "D:\\Caso\\Evidencia\\DDCapture.001" "E:\\Caso_DECIFRADO" <SHA-256 do relatório>')
item("Use uma pasta de saída FORA da pasta da evidência (ex.: pasta irmã com o sufixo _DECIFRADO).")
item("O processo lê a imagem inteira uma vez e mostra o progresso a cada 20 segundos (porcentagem, velocidade e "
     "tempo restante). Para 500 GB, conte algumas horas; o tempo depende sobretudo dos discos.")
item("Enquanto roda, o arquivo de saída tem a extensão .parcial. Só ao final ele vira .dd. Se o processo for "
     "interrompido, apague o .parcial e rode de novo.")
p("Ao terminar, a pasta de saída terá:")
item("<nome>_decifrado.dd: a imagem com o Windows decifrado;")
item("<nome>_decifrado_registro.txt: registro para a cadeia de custódia (hashes da origem e da imagem decifrada, "
     "partição, método, data, máquina e versões).")
aviso("confira no registro a linha \"Conferência: CONFERE\". Se aparecer \"NÃO CONFERE\", a imagem lida não é a "
      "mesma do relatório da extração: não prossiga sem esclarecer.")

h("8. PASSO 5 - Chave protegida por TPM ou senha")
p("Neste caso o disco só abre com a chave de recuperação (48 dígitos). Ela pode ser obtida:")
item("na conta Microsoft do usuário (requisição à Microsoft, informando o ID do protetor mostrado no PASSO 3);")
item("na TI do órgão, se o notebook for corporativo (domínio Active Directory ou Entra ID);")
item("na apreensão, com o notebook ligado (PASSO 0).")
p("Com a chave em mãos, rode o PASSO 4 acrescentando a chave ao final:")
cmd('decifrar.bat "D:\\Caso\\Evidencia\\DDCapture.001" "E:\\Caso_DECIFRADO" <SHA-256> '
    '123456-123456-123456-123456-123456-123456-123456-123456')

h("9. PASSO 6 - Processar no IPED e documentar")
item("Processe o arquivo <nome>_decifrado.dd no IPED como qualquer imagem de disco (mesmos parâmetros do setor). "
     "Agora o IPED verá o Windows, os usuários, navegadores etc.", True)
item("Guarde o _registro.txt junto com o resultado do IPED.", True)
item("No relatório, texto sugerido: \"A partição do Windows estava protegida por BitLocker com chave em claro, "
     "armazenada no próprio disco. O volume foi decifrado a partir da imagem, sem alteração do original, e a imagem "
     "decifrada (SHA-256 ...) foi processada no IPED.\" (Se foi usada chave de recuperação, dizer de onde ela veio.)",
     True)

h("10. Problemas comuns")
tabela(["Problema", "Solução"],
       [["\"Python não encontrado\" no instalar.bat", "Refaça o PASSO 1 marcando \"Add python.exe to PATH\"."],
        ["Imagem .E01 (EnCase)", "Converta para imagem bruta: FTK Imager > File > Export Disk Image > Raw (dd); "
                                 "confira o hash e use o .001 gerado."],
        ["\"NÃO CONFERE\" no registro", "SHA-256 informado errado ou imagem diferente da extração. Conferir antes "
                                       "de usar a imagem decifrada."],
        ["\"já existe; nada foi sobrescrito\"", "Já há uma imagem decifrada nessa pasta. Use outra pasta ou "
                                                "confirme que a existente é a correta."],
        ["Arquivo .parcial na saída", "Processo interrompido. Apague o .parcial e rode de novo."],
        ["Muito lento", "Grave a saída num SSD interno, não no mesmo HD USB da evidência."]], [6, 10.5])

h("11. Como funciona (para explicar no relatório ou em juízo)")
p("O BitLocker cifra o volume com uma chave (FVEK), que por sua vez é guardada protegida por \"protetores\" (TPM, "
  "senha, chave de recuperação). Quando o Windows liga a criptografia automaticamente e o usuário nunca entra com "
  "uma conta Microsoft ou corporativa, a chave fica gravada \"em claro\" no próprio disco, aguardando ativação: "
  "qualquer leitor de BitLocker abre o volume sem senha. Não há quebra de senha nem uso de chave externa.")
p("O kit usa a biblioteca forense de código aberto dissect (Fox-IT) para ler os metadados do BitLocker e obter a "
  "chave, e decifra o volume (AES-XTS) gravando uma nova imagem: cópia byte a byte da original, com a partição do "
  "Windows decifrada. A imagem original é aberta apenas para leitura. Qualquer perito pode repetir o procedimento com "
  "as mesmas ferramentas e deve chegar ao mesmo hash da imagem decifrada.")
p("Limitações: imagem bruta (.dd/.001); disco com tabela de partições GPT ou MBR; volume Windows (NTFS). BitLocker "
  "antigo (Windows 7/8, AES-CBC) funciona, porém mais devagar.")

d.save(os.path.join(AQUI, "TUTORIAL_BitLocker.docx"))
print("ok")
